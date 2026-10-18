"""Regenerates the synthetic DOCX/PDF/TXT fixtures used by tests/integration.test.mjs.

Run from the repo root: accessibility-parser/.venv/bin/python tests/fixtures/documents/make_fixtures.py
"""
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent


def png_bytes():
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    pix.clear_with(200)
    return pix.tobytes('png')


def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def course_docx():
    """Expected AA findings: 1.3.1 skipped heading, 1.1.1 missing alt, 2.4.4 'click here', 1.3.1 table headers."""
    doc = Document()
    doc.core_properties.title = 'ECE 30861 Syllabus'
    rpr = doc.styles['Normal'].element.get_or_add_rPr()
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'en-US')
    rpr.append(lang)
    image = OUT / '_tmp.png'
    image.write_bytes(png_bytes())
    doc.add_heading('ECE 30861 Syllabus', level=1)
    doc.add_paragraph('Welcome to the course.')
    doc.add_heading('Grading', level=3)
    doc.add_picture(str(image))
    described = doc.add_picture(str(image))
    image.unlink()
    pictures = list(doc.element.body.iter(qn('wp:docPr')))
    for pic in pictures:  # python-docx may write an empty descr; remove it so the first image has no alt text.
        pic.attrib.pop('descr', None)
    pictures[1].set('descr', 'Bar chart of the grade distribution')
    table = doc.add_table(rows=3, cols=2)
    for r, row in enumerate([('Item', 'Weight'), ('Exams', '60%'), ('Labs', '40%')]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    add_hyperlink(doc.add_paragraph('For the rubric, '), 'click here', 'https://example.edu/rubric')
    doc.save(OUT / 'course.docx')


def course_pdf():
    """Expected AA findings: 2.4.4 'click here' only. Image alt text is not extracted from PDFs."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), 'ECE 30861 Syllabus', fontsize=12)
    page.insert_text((72, 100), 'click here', fontsize=12)
    page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(70, 88, 140, 104), 'uri': 'https://example.edu/rubric'})
    page.insert_image(fitz.Rect(72, 120, 140, 188), stream=png_bytes())
    doc.set_metadata({'title': 'ECE 30861 Syllabus'})
    doc.xref_set_key(doc.pdf_catalog(), 'Lang', '(en-US)')
    doc.save(OUT / 'course.pdf')
    doc.close()


def encrypted_pdf():
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), 'Secret', fontsize=12)
    doc.save(OUT / 'encrypted.pdf', encryption=fitz.PDF_ENCRYPT_AES_256, user_pw='secret', owner_pw='owner')
    doc.close()


if __name__ == '__main__':
    course_docx()
    course_pdf()
    encrypted_pdf()
    (OUT / 'corrupt.docx').write_bytes(b'PK\x03\x04' + b'\x00not a real zip archive' * 4)
    (OUT / 'notes.txt').write_text('Week 1 notes\nRead chapter one before lab.\n')
    print('Fixtures written to', OUT)
