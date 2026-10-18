from pathlib import Path
import pytest
from docx import Document
from parser import parse_document, ParseError


def test_docx_structure(tmp_path):
    p=tmp_path/'syllabus.docx'; d=Document(); d.core_properties.title='Syllabus'
    d.add_heading('Course syllabus',level=1)
    d.add_paragraph('Course requirements')
    t=d.add_table(rows=2,cols=2); t.cell(0,0).text='Item'; t.cell(0,1).text='Points'
    t.cell(1,0).text='Quiz'; t.cell(1,1).text='10'
    d.save(p)
    parsed=parse_document(p)
    assert parsed['metadata']['title']=='Syllabus'
    assert parsed['elements'][0]['type']=='heading'
    assert parsed['elements'][0]['level']==1
    assert any(e['type']=='table' and e['rows']==2 for e in parsed['elements'])

def test_html_structure(tmp_path):
    p=tmp_path/'test.html'
    p.write_text('''<!doctype html><html lang="en"><head><title>Test</title></head><body>
<h1>Hello</h1><p style="color: #999999; font-size: 12pt">Intro <a href="https://example.org">click here</a></p>
<img src="x.png"><img src="dec.png" alt=""><table><tr><th>Title</th><th>Points</th></tr><tr><td>Quiz</td><td>10</td></tr></table>
</body></html>''')
    data=parse_document(p)
    assert data['metadata']['language']=='en'
    assert data['metadata']['title']=='Test'
    assert [e['type'] for e in data['elements']].count('image')==2
    assert [e['alt'] for e in data['elements'] if e['type']=='image']==[None,'']
    assert any(e['type']=='table' and e['has_header_row'] for e in data['elements'])
    assert any(e['type']=='link' and e['text']=='click here' for e in data['elements'])
    assert any(e['type']=='text' and e.get('color')=='#999999' for e in data['elements'])

def test_pdf(tmp_path):
    import pymupdf as fitz
    path=tmp_path/'page.pdf'; doc=fitz.open(); page=doc.new_page()
    page.insert_text((70,70),'Welcome to the course', fontsize=12)
    doc.set_metadata({'title':'Course'}); doc.save(path); doc.close()
    parsed=parse_document(path)
    assert parsed['metadata']['title']=='Course'
    assert any(e['type']=='text' and 'Welcome' in e['text'] for e in parsed['elements'])
    assert all(e['location']['page']==1 for e in parsed['elements'])

def test_txt_and_bad_file(tmp_path):
    p=tmp_path/'notes.txt'; p.write_text('Hello\n\nWorld\n')
    assert len(parse_document(p)['elements'])==2
    empty=tmp_path/'empty.txt';empty.touch()
    with pytest.raises(ParseError) as error:parse_document(empty)
    assert error.value.code=='UNPROCESSABLE_FILE'
    bad=tmp_path/'test.exe';bad.write_text('x')
    with pytest.raises(ParseError) as error:parse_document(bad)
    assert error.value.code=='UNSUPPORTED_FILE_TYPE'

FIXTURES=Path(__file__).resolve().parents[2]/'tests'/'fixtures'/'documents'

def test_docx_images_links_language_and_header_rows(tmp_path):
    parsed=parse_document(FIXTURES/'course.docx')
    assert parsed['metadata']['language']=='en-US'
    images=[e for e in parsed['elements'] if e['type']=='image']
    assert [i['alt'] for i in images]==[None,'Bar chart of the grade distribution']
    assert {'type':'link','text':'click here','href':'https://example.edu/rubric','location':{'line':7}} in parsed['elements']
    table=next(e for e in parsed['elements'] if e['type']=='table')
    assert table['has_header_row'] is False
    # Word's "Repeat as header row" marks a real header.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document(); t=d.add_table(rows=2,cols=1)
    tr_pr=t.rows[0]._tr.get_or_add_trPr(); tr_pr.append(OxmlElement('w:tblHeader'))
    p=tmp_path/'header.docx'; d.save(p)
    assert next(e for e in parse_document(p)['elements'] if e['type']=='table')['has_header_row'] is True

def test_unextractable_semantics_are_declared_for_the_engine():
    pdf=parse_document(FIXTURES/'course.pdf')
    assert pdf['metadata']['not_extracted']==['alt_text','headings','tables']
    assert pdf['metadata']['language']=='en-US'
    txt=parse_document(FIXTURES/'notes.txt')
    assert set(txt['metadata']['not_extracted'])>={'title','language','headings'}
    assert 'not_extracted' not in parse_document(FIXTURES/'course.docx')['metadata']

def test_corrupt_and_encrypted_documents_raise_parse_errors():
    for name in ['corrupt.docx','encrypted.pdf']:
        with pytest.raises(ParseError) as error: parse_document(FIXTURES/name)
        assert error.value.code=='UNPROCESSABLE_FILE'

def test_cli_stdout_is_only_json():
    import json, subprocess, sys
    root=Path(__file__).resolve().parents[1]
    for name in ['course.pdf','corrupt.docx']:
        out=subprocess.run([sys.executable,'-m','parser.cli',str(FIXTURES/name)],cwd=root,capture_output=True,text=True).stdout
        json.loads(out)  # A library warning on stdout would make this fail.
